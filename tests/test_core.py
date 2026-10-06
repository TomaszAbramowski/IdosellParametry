"""Testy na atrapach - bez połączenia z IdoSellem. Uruchom: python -m unittest -v"""
import os
import tempfile
import unittest
from pathlib import Path
from functools import cmp_to_key
from unittest.mock import MagicMock, patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")   # Qt bez otwierania okien

from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication

from app.dictionary import ParameterDictionary, PolishState, TranslationState, evaluate_polish, evaluate_translation
from app.idosell_api import FetchCancelled, IdoSellClient
from app.main_window import parameter_columns
from app.natural_sort import natural_compare
from app.products_tab import summarize_parameters
from app.table_model import Column


def value_columns(dictionary, products):
    """Proste kolumny wartości do testów filtrów tabeli."""
    return [Column("Parametr", lambda e: dictionary.parent_name(e)), Column("Wartość", lambda e: e.name_pl),
            Column("EN", lambda e: e.name_en), Column("Tłumaczenie EN")]
from app.products import Product, ProductIndex
from app.table_model import STATE_FILTERS, DictionaryFilterProxy, DictionaryTableModel


def raw(element_id, element_type, polish, english, parameter_id=0, value_ids=()):
    return {"id": element_id, "type": element_type, "parameterId": str(parameter_id),
            "names": [{"languageId": "pol", "value": polish}, {"languageId": "eng", "value": english}],
            "parameterValueIds": list(value_ids)}


SAMPLE = {
    10502: raw(10502, "parameter", "Bęben", "Bęben", value_ids=[10565, 10566]),
    10392: raw(10392, "parameter", "Materiał rękojeści", "Handle material", value_ids=[10604]),
    10565: raw(10565, "value", "5-cio strzałowy", "5-cio strzałowy", parameter_id=10502),
    10566: raw(10566, "value", "10", "10", parameter_id=10502),
    10604: raw(10604, "value", "drewno", "", parameter_id=10392),
    10700: raw(10700, "value", "dane do uzupełnienia", "dane do uzupełnienia", parameter_id=10392),
    10701: raw(10701, "value", "brak", "brak", parameter_id=10502),
}


def raw_product(product_id, name, category, parameters):
    """Kształt jak w odpowiedzi IdoSella: parameters = {ID parametru: [ID wartości, ...]}."""
    return {"productId": product_id, "productDisplayedCode": f"KOD-{product_id}", "categoryName": category,
            "productDescriptionsLangData": [{"langId": "eng", "productName": f"EN {name}"},
                                            {"langId": "pol", "productName": name}],
            "productParameters": [{"parameterId": parameter_id, "parameterType": "parameter",
                                   "parameterValues": [{"parameterValueId": v} for v in values]}
                                  for parameter_id, values in parameters.items()]}


PRODUCTS = [
    Product.from_api(raw_product(1, "Rewolwer B", "Rewolwery", {10502: [10565]})),
    Product.from_api(raw_product(2, "Rewolwer A", "Rewolwery", {10502: [10565, 10565]})),
    Product.from_api(raw_product(3, "Nóż", "Noże", {10392: [10604]})),
]


class FakeResponse:
    def __init__(self, status_code, data=None, text=""):
        self.status_code, self._data, self.text = status_code, data, text

    def json(self):
        return self._data


def page(number, count, elements):
    return FakeResponse(200, {"parametersResult": {str(e["id"]): e for e in elements},
                              "resultsNumberPage": count, "resultsNumberAll": 99})


class EvaluateTranslationTests(unittest.TestCase):
    def test_states(self):
        self.assertIs(evaluate_translation("Nie", ""), TranslationState.MISSING)
        self.assertIs(evaluate_translation("Nie", "Nie"), TranslationState.SAME_AS_POLISH)
        self.assertIs(evaluate_translation("129", "129"), TranslationState.NUMBER_OR_CODE)
        self.assertIs(evaluate_translation("44.2 J / 32.6 ft.lbs", "44.2 J / 32.6 ft.lbs"), TranslationState.NUMBER_OR_CODE)
        self.assertIs(evaluate_translation("5.5 mm (.22)", "5.5 mm (.22)"), TranslationState.NUMBER_OR_CODE)
        self.assertIs(evaluate_translation("6-cio strzałowy", "6-cio strzałowy"), TranslationState.SAME_AS_POLISH)
        self.assertIs(evaluate_translation("10 szt.", "10 szt."), TranslationState.SAME_AS_POLISH)
        self.assertIs(evaluate_translation("Nie", "No"), TranslationState.TRANSLATED)
        self.assertIs(evaluate_translation("dane do uzupełnienia", "x"), TranslationState.PLACEHOLDER)


class EvaluatePolishTests(unittest.TestCase):
    def test_states(self):
        self.assertIs(evaluate_polish("drewno"), PolishState.OK)
        self.assertIs(evaluate_polish(""), PolishState.EMPTY)
        self.assertIs(evaluate_polish(" - "), PolishState.EMPTY)
        self.assertIs(evaluate_polish("Dane do uzupełnienia"), PolishState.PLACEHOLDER)
        self.assertIs(evaluate_polish("b/d"), PolishState.PLACEHOLDER)
        self.assertIs(evaluate_polish("Brak"), PolishState.NONE_VALUE)
        self.assertIs(evaluate_polish("129"), PolishState.OK)


class ParameterDictionaryTests(unittest.TestCase):
    def test_split_and_counts(self):
        dictionary = ParameterDictionary(SAMPLE)
        self.assertEqual([p.id for p in dictionary.parameters], [10502, 10392])   # Bęben, Materiał
        self.assertEqual(len(dictionary.values), 5)
        self.assertEqual(dictionary.parent_name(dictionary.elements[10565]), "Bęben")
        self.assertEqual(dictionary.untranslated_values_of(10502), 2)   # 5-cio strzałowy + „brak”; „10” to liczba
        self.assertEqual(dictionary.untranslated_values_of(10392), 1)   # drewno; zaślepki nie tłumaczymy
        self.assertEqual(dictionary.polish_problem_values_of(10392), 1)

    def test_string_keys_from_json_cache(self):
        dictionary = ParameterDictionary({str(k): v for k, v in SAMPLE.items()})
        self.assertIn(10565, dictionary.elements)


class ProductTests(unittest.TestCase):
    def test_from_api(self):
        product = PRODUCTS[1]
        self.assertEqual((product.id, product.code, product.name, product.category), (2, "KOD-2", "Rewolwer A", "Rewolwery"))
        self.assertEqual(product.value_ids, (10565,))      # powtórzenia usunięte
        self.assertEqual(product.parameter_ids, (10502,))
        self.assertEqual(Product.from_cache(product.to_cache()), product)

    def test_index(self):
        index = ProductIndex(PRODUCTS)
        self.assertEqual([p.name for p in index.products_with(10565)], ["Rewolwer A", "Rewolwer B"])
        self.assertEqual(index.categories_of(10565), ("Rewolwery",))
        self.assertEqual(index.products_with(99999), [])
        self.assertEqual(index.product_count_with_parameter(10502), 2)
        self.assertEqual(index.categories, ["Noże", "Rewolwery"])
        self.assertEqual(index.used_value_ids, {10565, 10604})

    def test_products_with_all(self):
        products = [*PRODUCTS, Product.from_api(raw_product(4, "Rewolwer C", "Rewolwery", {10502: [10565, 10701]}))]
        index = ProductIndex(products)
        self.assertEqual([p.id for p in index.products_with_all([10565])], [2, 1, 4])
        self.assertEqual([p.id for p in index.products_with_all([10565, 10701])], [4])   # tylko C ma obie
        self.assertEqual(index.products_with_all([10565, 10604]), [])
        self.assertEqual(index.products_with_all([]), [])


@patch("app.idosell_api.time.sleep")   # testy bez czekania
class IdoSellClientTests(unittest.TestCase):
    def client_with(self, responses):
        client = IdoSellClient("https://sklep.example", "klucz")
        client._session = MagicMock()
        client._session.request.side_effect = responses
        return client

    def test_fetches_all_pages(self, _sleep):
        client = self.client_with([page(0, 2, [SAMPLE[10502]]), page(1, 2, [SAMPLE[10565]])])
        progress = []
        elements = client.fetch_whole_dictionary(on_page=lambda done, total: progress.append((done, total)))
        self.assertEqual(set(elements), {10502, 10565})
        self.assertEqual(progress, [(1, 2), (2, 2)])
        body = client._session.request.call_args_list[1].kwargs["json"]
        self.assertEqual(body["params"]["resultsPage"], 1)
        self.assertEqual(body["params"]["languagesIds"], ["pol", "eng"])

    def test_fetches_products(self, _sleep):
        responses = [FakeResponse(200, {"results": [raw_product(1, "A", "K", {})], "resultsNumberPage": 2}),
                     FakeResponse(200, {"results": [raw_product(2, "B", "K", {})], "resultsNumberPage": 2})]
        client = self.client_with(responses)
        self.assertEqual([p["productId"] for p in client.fetch_all_products()], [1, 2])
        call = client._session.request.call_args_list[0]
        self.assertEqual(call.args[:2], ("POST", "https://sklep.example/api/admin/v8/products/products/search"))
        self.assertIn("parameters", call.kwargs["json"]["params"]["returnElements"])

    def test_products_limit_small(self, _sleep):
        # Serwer ma 5 stron, ale chcemy tylko 20 -> jedno zapytanie po 20 sztuk
        client = self.client_with([FakeResponse(200, {"results": [raw_product(i, "P", "K", {}) for i in range(20)],
                                                      "resultsNumberPage": 5})])
        progress = []
        products = client.fetch_all_products(on_page=lambda d, t: progress.append((d, t)), limit=20)
        self.assertEqual(len(products), 20)
        self.assertEqual(client._session.request.call_count, 1)
        self.assertEqual(client._session.request.call_args.kwargs["json"]["params"]["resultsLimit"], 20)
        self.assertEqual(progress, [(1, 1)])

    def test_products_limit_over_one_page(self, _sleep):
        pages = [FakeResponse(200, {"results": [raw_product(n * 100 + i, "P", "K", {}) for i in range(100)],
                                    "resultsNumberPage": 9}) for n in range(2)]
        client = self.client_with(pages)
        self.assertEqual(len(client.fetch_all_products(limit=150)), 150)   # 2 strony po 100, przycięte do 150
        self.assertEqual(client._session.request.call_count, 2)

    def test_guard_blocks_write(self, _sleep):
        client = self.client_with([])
        with self.assertRaises(PermissionError):
            client._request("PUT", "products/parameters", {})
        with self.assertRaises(PermissionError):
            client._request("POST", "products/parameters/delete", {})
        client._session.request.assert_not_called()

    def test_retries_server_error(self, _sleep):
        client = self.client_with([FakeResponse(503, text="busy"), page(0, 1, [SAMPLE[10502]])])
        self.assertEqual(set(client.fetch_whole_dictionary()), {10502})

    def test_no_retry_on_auth_error(self, _sleep):
        client = self.client_with([FakeResponse(401, text="zły klucz")])
        with self.assertRaises(RuntimeError):
            client.fetch_whole_dictionary()
        self.assertEqual(client._session.request.call_count, 1)

    def test_cancel(self, _sleep):
        client = self.client_with([])
        with self.assertRaises(FetchCancelled):
            client.fetch_whole_dictionary(should_stop=lambda: True)

    def test_missing_settings(self, _sleep):
        with self.assertRaises(ValueError):
            IdoSellClient("", "")


class FilterTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.application = QApplication.instance() or QApplication([])

    def setUp(self):
        self.dictionary = ParameterDictionary(SAMPLE)
        self.drafts = {}
        self.model = DictionaryTableModel(self.drafts, lambda: None)
        self.model.set_content(self.dictionary.values, value_columns(self.dictionary, None))
        self.proxy = DictionaryFilterProxy()
        self.proxy.setSourceModel(self.model)

    def visible_ids(self):
        return {self.model.element_at(self.proxy.mapToSource(self.proxy.index(r, 0)).row()).id
                for r in range(self.proxy.rowCount())}

    def test_untranslated_of_one_parameter(self):
        self.proxy.set_filters("", STATE_FILTERS[1], 10502)
        self.assertEqual(self.visible_ids(), {10565, 10701})

    def test_polish_filters(self):
        labels = [f.label for f in STATE_FILTERS]
        self.proxy.set_filters("", STATE_FILTERS[labels.index("Braki po polsku (pusta / dane do uzupełnienia)")], None)
        self.assertEqual(self.visible_ids(), {10700})
        self.proxy.set_filters("", STATE_FILTERS[labels.index("Po polsku „brak” – do sprawdzenia")], None)
        self.assertEqual(self.visible_ids(), {10701})

    def test_text_search_and_drafts(self):
        translation_column = len(value_columns(self.dictionary, None)) - 1
        row = next(r for r in range(self.model.rowCount()) if self.model.element_at(r).id == 10565)
        self.assertTrue(self.model.setData(self.model.index(row, translation_column), " 5-shot "))
        self.assertEqual(self.drafts, {10565: "5-shot"})
        self.proxy.set_filters("shot", STATE_FILTERS[0], None)
        self.assertEqual(self.visible_ids(), {10565})
        self.proxy.set_filters("", STATE_FILTERS[2], None)            # do tłumaczenia, bez mojego
        self.assertEqual(self.visible_ids(), {10604, 10701})
        self.model.setData(self.model.index(row, translation_column), "")   # puste = usuń
        self.assertEqual(self.drafts, {})

    def test_category_filter_and_product_search(self):
        index = ProductIndex(PRODUCTS)
        self.model.set_content(self.dictionary.values, value_columns(self.dictionary, index),
                               lambda e: " ".join(p.name for p in index.products_with(e.id)))
        self.proxy.category_lookup = index.categories_of
        self.proxy.set_filters("", STATE_FILTERS[0], None, "Noże")
        self.assertEqual(self.visible_ids(), {10604})
        self.proxy.set_filters("rewolwer a", STATE_FILTERS[0], None, None)
        self.assertEqual(self.visible_ids(), {10565})

    def test_only_used_filter(self):
        self.proxy.used_value_ids = ProductIndex(PRODUCTS).used_value_ids
        self.proxy.set_filters("", STATE_FILTERS[0], None, None, only_used=True)
        self.assertEqual(self.visible_ids(), {10565, 10604})

    def test_parameter_columns_build(self):
        model = DictionaryTableModel({}, lambda: None)
        model.set_content(self.dictionary.parameters, parameter_columns(self.dictionary, None))
        self.assertEqual(model.rowCount(), 2)


class NaturalSortTests(unittest.TestCase):
    def test_numbers(self):
        names = ["130", "14", "2", "150", "16"]
        self.assertEqual(sorted(names, key=cmp_to_key(natural_compare)), ["2", "14", "16", "130", "150"])

    def test_text_and_units(self):
        self.assertLess(natural_compare("5-cio strzałowy", "6-cio strzałowy"), 0)
        self.assertLess(natural_compare("200 bar", "1000 bar"), 0)
        self.assertLess(natural_compare("Sakwa", "Ściereczka"), 0)     # polska kolejność: s < ś
        self.assertLess(natural_compare("Ściereczka", "Torba"), 0)


class ValuesTabTests(unittest.TestCase):
    """Tryb 2: tylko nazwy - parametry po lewej, wartości wybranego parametru po prawej, edycja jak w trybie 1."""

    @classmethod
    def setUpClass(cls):
        cls.application = QApplication.instance() or QApplication([])

    def setUp(self):
        from app.values_tab import ParameterValuesTab
        self.dictionary = ParameterDictionary(SAMPLE)
        self.drafts, self.changes = {}, []
        self.tab = ParameterValuesTab(self.drafts, lambda: self.changes.append(1))
        self.tab.set_content(self.dictionary, ProductIndex(PRODUCTS))

    def visible(self):
        return [v.id for v in self.tab.visible_values()]

    def test_only_name_columns(self):
        headers = [self.tab.values.proxy.headerData(c, Qt.Orientation.Horizontal)
                   for c in range(self.tab.values.proxy.columnCount())]
        self.assertEqual(headers, ["Wartość po polsku", "Po angielsku – teraz w sklepie", "Po angielsku – NOWE ✎"])
        list_headers = [self.tab._list.horizontalHeaderItem(c).text() for c in range(self.tab._list.columnCount())]
        self.assertEqual(list_headers, ["Parametr po polsku", "Po angielsku – teraz w sklepie", "Po angielsku – NOWE ✎"])

    def test_select_parameter_shows_only_its_values_in_natural_order(self):
        self.tab.select(10502)
        self.assertEqual(self.visible(), [10565, 10566, 10701])          # 5-cio…, 10, brak
        self.assertEqual(len(self.tab.all_filtered_values()), 5)          # wszystkie parametry
        self.assertEqual(self.visible(), [10565, 10566, 10701])           # wybór parametru zostaje

    def test_edit_value_translation_like_mode_1(self):
        self.tab.select(10502)
        proxy = self.tab.values.proxy
        row = next(r for r in range(proxy.rowCount()) if self.tab.values.element_at(proxy.index(r, 0)).id == 10565)
        proxy.setData(proxy.index(row, proxy.columnCount() - 1), "5-shot")
        self.assertEqual(self.drafts, {10565: "5-shot"})
        self.assertTrue(self.changes)

    def test_edit_parameter_translation_in_list(self):
        from app.values_tab import NAME, TRANSLATION
        row = next(r for r in range(self.tab._list.rowCount())
                   if self.tab._list.item(r, NAME).data(Qt.ItemDataRole.UserRole) == 10502)
        self.tab._list.item(row, TRANSLATION).setText(" Cylinder ")
        self.assertEqual(self.drafts, {10502: "Cylinder"})
        self.tab._list.item(row, TRANSLATION).setText("")
        self.assertEqual(self.drafts, {})

    def test_reveal_value_selects_its_parameter(self):
        self.tab.reveal(10604)
        self.assertEqual(self.tab._current.id, 10392)
        self.assertEqual(self.visible(), [10700, 10604])

    def test_only_todo_hides_done_parameters(self):
        from app.values_tab import NAME
        self.drafts.update({10565: "x", 10701: "y", 10604: "z", 10502: "Cylinder"})
        self.tab._refresh_now()
        self.tab._only_todo_check.setChecked(True)
        hidden = {self.tab._list.item(r, NAME).data(Qt.ItemDataRole.UserRole): self.tab._list.isRowHidden(r)
                  for r in range(1, self.tab._list.rowCount())}
        self.assertEqual(hidden, {10502: True, 10392: True})


class SummaryTests(unittest.TestCase):
    def test_common_and_different(self):
        dictionary = ParameterDictionary(SAMPLE)
        products = [Product.from_api(raw_product(1, "A", "K", {10502: [10565], 10392: [10604]})),
                    Product.from_api(raw_product(2, "B", "K", {10502: [10701]}))]
        summaries = {s.parameter.id: s for s in summarize_parameters(products, dictionary)}
        drum = summaries[10502]
        self.assertEqual(drum.product_count, 2)
        self.assertEqual({v.id for v, _ in drum.values}, {10565, 10701})        # różne wartości
        self.assertEqual(summaries[10392].product_count, 1)                    # tylko produkt A


class ExcelTests(unittest.TestCase):
    """Zgranie -> „ręczna” edycja Excela -> wgranie: do podglądu trafiają tylko różnice. Osobno tryb 1 i 2."""

    def setUp(self):
        from openpyxl import load_workbook
        import app.excel_io as excel_io
        self.load_workbook = load_workbook
        folder = Path(tempfile.mkdtemp())
        patcher = patch.object(excel_io, "DATA_DIRECTORY", folder)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.excel_io = excel_io
        self.dictionary = ParameterDictionary(SAMPLE)
        drafts = {10565: "5-shot"}
        self.parameters_path = excel_io.export_parameters(self.dictionary, drafts, self.dictionary.parameters)
        self.values_path = excel_io.export_values(self.dictionary, drafts, self.dictionary.values)

    def edit(self, path, sheet_name, edits):
        workbook = self.load_workbook(path)
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[2]]                 # wiersz 1 = opis stref
        id_column, new_column = headers.index("Nr (ID)") + 1, headers.index("Po angielsku – NOWE ✎") + 1
        for row in range(3, sheet.max_row + 1):
            if sheet.cell(row, id_column).value in edits:
                sheet.cell(row, new_column).value = edits[sheet.cell(row, id_column).value]
        workbook.save(path)

    def read(self, path, sheet):
        return self.excel_io.read_changes_from_excel(path, self.dictionary, sheet)

    def test_separate_files(self):
        parameters = self.load_workbook(self.parameters_path).sheetnames
        values = self.load_workbook(self.values_path).sheetnames
        self.assertEqual(parameters, ["Instrukcja", "Parametry"])
        self.assertEqual(values, ["Instrukcja", "Wartości"])

    def test_draft_exported_and_counted_as_change(self):
        result = self.read(self.values_path, "Wartości")
        self.assertEqual([(c.element_id, c.old_en, c.new_en) for c in result.changes],
                         [(10565, "5-cio strzałowy", "5-shot")])
        self.assertEqual(self.read(self.parameters_path, "Parametry").changes, [])

    def test_only_differences(self):
        self.edit(self.parameters_path, "Parametry", {10502: "Cylinder", 10392: "Handle material"})  # 10392 = jak w sklepie
        result = self.read(self.parameters_path, "Parametry")
        self.assertEqual({c.element_id: c.new_en for c in result.changes}, {10502: "Cylinder"})
        self.assertEqual(result.unchanged, 1)
        self.edit(self.values_path, "Wartości", {10566: 10, 10604: " wood "})       # 10 jako liczba = jak w sklepie
        result = self.read(self.values_path, "Wartości")
        self.assertEqual({c.element_id: c.new_en for c in result.changes}, {10565: "5-shot", 10604: "wood"})
        self.assertEqual(next(c for c in result.changes if c.element_id == 10604).parameter, "Materiał rękojeści")

    def test_legacy_file_still_readable(self):
        """Starsze pliki: nagłówki w 1. wierszu i dawne nazwy kolumn."""
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Parametry"
        sheet.append(["ID", "Nazwa PL", "EN w sklepie", "Nowa nazwa EN"])
        sheet.append([10502, "Bęben", "Bęben", "Cylinder"])
        path = Path(tempfile.mkdtemp()) / "stary.xlsx"
        workbook.save(path)
        self.assertEqual([c.new_en for c in self.read(path, "Parametry").changes], ["Cylinder"])

    def test_wrong_file_for_mode(self):
        with self.assertRaises(self.excel_io.WrongFileError) as error:
            self.read(self.values_path, "Parametry")
        self.assertIn("drugiej zakładce", str(error.exception))

    def test_unknown_id_and_wrong_type_are_reported(self):
        workbook = self.load_workbook(self.values_path)
        sheet = workbook["Wartości"]
        headers = [cell.value for cell in sheet[2]]
        for element_id in (99999, 10502):                 # nieistniejące ID i parametr w arkuszu wartości
            row = [None] * len(headers)
            row[headers.index("Nr (ID)")], row[headers.index("Po angielsku – NOWE ✎")] = element_id, "x"
            sheet.append(row)
        workbook.save(self.values_path)
        problems = self.read(self.values_path, "Wartości").problems
        self.assertTrue(any("99999" in p for p in problems))
        self.assertTrue(any("10502" in p and "nie pasuje" in p for p in problems))


class FakeShop:
    """Udaje IdoSella: słownik nazw w pamięci; przyjmuje tylko jeden „kształt” zapisu, inne odrzuca kodem 422."""

    def __init__(self, names: dict[int, dict[str, str]], accepted_shape: str):
        self.names, self.accepted_shape, self.writes = names, accepted_shape, []

    def read_elements(self, ids):
        return {i: {"id": i, "names": [{"languageId": l, "value": v} for l, v in self.names[i].items()]}
                for i in ids if i in self.names}

    def write_names(self, items, shape):
        from app.idosell_api import ensure_only_names
        ensure_only_names(items)
        self.writes.append(shape)
        if shape != self.accepted_shape:
            return 422, "zły format"
        for item in items:
            self.names[item["id"]] = {n["lang_id"]: n["value"] for n in item["names"]}   # zastępuje wszystkie
        return 200, "{}"


class SenderTests(unittest.TestCase):
    def setUp(self):
        import app.sender as sender
        self.sender = sender
        patcher = patch.object(sender, "BACKUP_DIRECTORY", Path(tempfile.mkdtemp()))
        patcher.start()
        self.addCleanup(patcher.stop)
        self.dictionary = ParameterDictionary(SAMPLE)

    def shop(self, accepted_shape="plain"):
        return FakeShop({10565: {"pol": "5-cio strzałowy", "eng": "5-cio strzałowy", "ger": "5-cio strzałowy"},
                         10502: {"pol": "Bęben", "eng": "Bęben", "ger": "Trommel"}}, accepted_shape)

    def changes(self):
        from app.changes import make_change
        return [make_change(self.dictionary.elements[10565], self.dictionary, "5-shot"),
                make_change(self.dictionary.elements[10502], self.dictionary, "Cylinder")]

    def test_first_send_is_single_test(self):
        shop = self.shop("plain")
        report = self.sender.send_changes(shop, self.changes(), known_shape=None)
        self.assertTrue(report.test_mode)
        self.assertEqual(report.count(self.sender.OK), 1)               # tylko jedna zmiana
        self.assertEqual(report.shape, "plain")                         # format z dokumentacji
        self.assertEqual(shop.writes, ["plain"])
        self.assertEqual(shop.names[10565], {"pol": "5-cio strzałowy", "eng": "5-shot", "ger": "5-cio strzałowy"})
        self.assertEqual(shop.names[10502]["eng"], "Bęben")              # drugiej nie ruszono
        self.assertTrue(report.backup_path.exists())

    def test_normal_send_and_skip_already_done(self):
        shop = self.shop("plain")
        shop.names[10502]["eng"] = "Cylinder"                           # ktoś już to zrobił w panelu
        report = self.sender.send_changes(shop, self.changes(), known_shape="plain")
        statuses = {o.change.element_id: o.status for o in report.outcomes}
        self.assertEqual(statuses, {10565: self.sender.OK, 10502: self.sender.SKIPPED})
        self.assertEqual(shop.names[10502]["ger"], "Trommel")

    def test_conflict_when_shop_changed(self):
        shop = self.shop("plain")
        shop.names[10565]["eng"] = "five shots"                         # zmienione w panelu po pobraniu słownika
        report = self.sender.send_changes(shop, self.changes()[:1], known_shape="plain")
        self.assertEqual(report.outcomes[0].status, self.sender.SKIPPED)
        self.assertEqual(shop.names[10565]["eng"], "five shots")        # nie nadpisane

    def test_verification_catches_lost_language_and_restore(self):
        shop = self.shop("plain")
        original_write = shop.write_names
        def broken_write(items, shape):                                  # sklep „gubi” język niemiecki
            status, text = original_write(items, shape)
            for item in items:
                shop.names[item["id"]].pop("ger", None)
            return status, text
        shop.write_names = broken_write
        report = self.sender.send_changes(shop, self.changes()[:1], known_shape="plain")
        self.assertEqual(report.outcomes[0].status, self.sender.FAILED)
        self.assertIn("ger", report.outcomes[0].message)
        shop.write_names = original_write
        restored, failed = self.sender.restore_from_backup(shop, report.outcomes, report.last_shape)
        self.assertEqual((restored, failed), (1, 0))
        self.assertEqual(shop.names[10565]["eng"], "5-cio strzałowy")


class ResponseFaultTests(unittest.TestCase):
    def test_faults_are_read(self):
        from app.sender import response_faults
        text = '{"results": [{"item_id": 5, "faultCode": 0}, {"item_id": 7, "faultCode": 3, "faultString": "zły"}],' \
               ' "errors": {"faultCode": 1, "faultString": "ogólny"}}'
        self.assertEqual(response_faults(text), {7: "3: zły", 0: "1: ogólny"})
        self.assertEqual(response_faults("nie json"), {})


class GuardTests(unittest.TestCase):
    def test_only_names_allowed(self):
        from app.idosell_api import ensure_only_names
        ensure_only_names([{"id": 1, "names": [{"lang_id": "eng", "value": "x"}]}])
        for bad in ([], [{"id": 1, "names": [], }], [{"id": 1, "names": [{"lang_id": "eng", "value": "x"}],
                                                      "descriptions": []}],
                    [{"id": 1, "names": [{"lang_id": "eng", "value": "x", "extra": 1}]}],
                    [{"id": 1, "names": [{"languageId": "eng", "value": "x"}]}],      # pole z odczytu - zły klucz
                    [{"id": i, "names": [{"lang_id": "eng", "value": "x"}]} for i in range(21)]):
            with self.assertRaises(PermissionError):
                ensure_only_names(bad)

    @patch("app.idosell_api.time.sleep")
    def test_delete_still_blocked(self, _sleep):
        client = IdoSellClient("https://sklep.example", "klucz")
        client._session = MagicMock()
        with self.assertRaises(PermissionError):
            client._request("POST", "products/parameters/delete", {})
        with self.assertRaises(PermissionError):
            client._request("PUT", "products/parameters", {})              # zapis tylko przez write_names
        client._session.request.assert_not_called()


if __name__ == "__main__":
    unittest.main()
